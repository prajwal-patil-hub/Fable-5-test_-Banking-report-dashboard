"""End-to-end: API surface against a seeded app, plus the full
document-ingestion round trip using a synthetic annual-report PDF."""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))


def _demo_bank_id(client) -> int:
    banks = client.get("/api/banks").json()["banks"]
    return next(b["id"] for b in banks if b["is_demo"])


def test_banks_years_kpis(client):
    banks = client.get("/api/banks").json()["banks"]
    assert len(banks) >= 40  # Indian roster + demo institutions
    demo = [b for b in banks if b["is_demo"]]
    assert len(demo) == 4 and all(b["has_data"] for b in demo)

    bank_id = demo[0]["id"]
    years = client.get(f"/api/banks/{bank_id}/years").json()["fiscal_years"]
    assert years == ["FY2023", "FY2024", "FY2025"]

    payload = client.get(f"/api/banks/{bank_id}/kpis", params={"fiscal_year": "FY2025"}).json()
    kpis = {k["kpi_code"]: k for k in payload["kpis"]}
    assert {"roe", "gnpa_ratio", "crar", "casa_ratio"} <= set(kpis)
    roe = kpis["roe"]
    assert roe["lineage"]["extraction_method"] in ("derived", "manual")
    assert roe["yoy_change"] is not None  # FY2024 exists
    assert roe["validation_status"] in ("passed", "warning", "failed")


def test_history_benchmarking_narrative_validations(client):
    bank_id = _demo_bank_id(client)

    hist = client.get(f"/api/banks/{bank_id}/kpis/pat/history").json()
    assert [p["fiscal_year"] for p in hist["series"]] == ["FY2023", "FY2024", "FY2025"]

    bench = client.get("/api/benchmarking",
                       params={"kpi_code": "roe", "fiscal_year": "FY2025"}).json()
    assert len(bench["peers"]) == 4

    narrative = client.get(f"/api/banks/{bank_id}/narrative",
                           params={"fiscal_year": "FY2025"}).json()
    assert narrative["executive_summary"]

    validations = client.get(f"/api/banks/{bank_id}/validations",
                             params={"fiscal_year": "FY2025"}).json()["results"]
    assert validations and all(v["status"] in ("passed", "failed") for v in validations)


def test_unknown_bank_and_kpi_404(client):
    assert client.get("/api/banks/9999/years").status_code == 404
    assert client.get("/api/benchmarking",
                      params={"kpi_code": "nope", "fiscal_year": "FY2025"}).status_code == 404


def test_document_upload_roundtrip(client, tmp_path):
    """The flagship E2E: synthetic annual report PDF → extraction → derivation
    → validation, all through the public API."""
    import make_sample_pdf

    pdf_path = tmp_path / "pinnacle_fy2025.pdf"
    make_sample_pdf.build(str(pdf_path))

    bank_id = _demo_bank_id(client)
    resp = client.post(
        "/api/documents/upload",
        data={"bank_id": bank_id, "doc_type": "annual_report", "fiscal_year": "FY2026"},
        files={"file": ("pinnacle_fy2025.pdf", pdf_path.read_bytes(), "application/pdf")},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["status"] == "processed"
    assert body["extracted_count"] >= 12  # table + narrative figures

    kpis = client.get(f"/api/banks/{bank_id}/kpis",
                      params={"fiscal_year": "FY2026"}).json()["kpis"]
    by_code = {k["kpi_code"]: k for k in kpis}
    # Table figures land with current-year column values
    assert by_code["gnpa"]["value"] == 5400.0
    assert by_code["cet1_ratio"]["value"] == 14.2
    # Narrative-text figures are captured by the rule-based extractor
    assert by_code["nim"]["value"] == 3.95
    # Ratios not in the document are derived from extracted components
    assert by_code["casa_ratio"]["lineage"]["extraction_method"] == "derived"
    assert by_code["casa_ratio"]["value"] == pytest.approx(40.7, abs=0.2)
    # Lineage points back at the uploaded document
    assert by_code["gnpa"]["lineage"]["document_id"] == body["document_id"]
    assert by_code["gnpa"]["lineage"]["page"] is not None

    # The document titles itself FY2025 but was filed under FY2026 —
    # the fiscal-year cross-check must flag it
    validations = client.get(f"/api/banks/{bank_id}/validations",
                             params={"fiscal_year": "FY2026"}).json()["results"]
    crosscheck = next(v for v in validations if v["rule_code"] == "fy_crosscheck")
    assert crosscheck["status"] == "failed"
    assert "FY2025" in crosscheck["message"]


def test_exports_produce_valid_files(client):
    bank_id = _demo_bank_id(client)
    params = {"bank_id": bank_id, "fiscal_year": "FY2025"}

    excel = client.get("/api/exports/excel", params=params)
    assert excel.status_code == 200 and excel.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(excel.content))
    assert "Validated KPIs" in wb.sheetnames and "Benchmarking" in wb.sheetnames

    pptx = client.get("/api/exports/pptx", params=params)
    assert pptx.status_code == 200 and pptx.content[:2] == b"PK"
    from pptx import Presentation
    deck = Presentation(io.BytesIO(pptx.content))
    assert sum(1 for _ in deck.slides) >= 8

    pdf = client.get("/api/exports/pdf", params=params)
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
