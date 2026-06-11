"""Excel Data Pack — full lineage preserved: every KPI row carries source
document, page, method and confidence so the workbook is auditable standalone."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import Bank
from app.modules.benchmarking.engine import benchmark_summary
from app.modules.export.theme import CATEGORY_LABELS, GOLD, INK, TEXT_PRIMARY
from app.modules.kpi_warehouse import service as warehouse

HEADER_FILL = PatternFill("solid", fgColor=INK)
HEADER_FONT = Font(bold=True, color=GOLD)
TITLE_FONT = Font(bold=True, size=14, color=INK)


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for col_idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[col_idx - 1])) for r in rows if r[col_idx - 1] is not None] or [10])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 3, 60)
    ws.freeze_panes = "A2"


def build_excel(db: Session, bank: Bank, fy: str) -> bytes:
    payload = warehouse.kpi_payload(db, bank, fy)
    kpis = payload["kpis"]

    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["B2"] = "SOVEREIGN — Banking Intelligence Data Pack"
    cover["B2"].font = TITLE_FONT
    cover["B4"] = f"Institution: {bank.name} ({bank.code})"
    cover["B5"] = f"Fiscal Year: {fy}"
    cover["B6"] = "All monetary figures in ₹ crore. Single source: Sovereign KPI warehouse."
    cover.column_dimensions["B"].width = 80

    kpi_headers = ["KPI", "Category", "Value", "Unit", "YoY Δ", "YoY Δ%",
                   "Validation", "Method", "Confidence", "Source Document", "Page", "Source Text"]

    def kpi_row(k: dict) -> list:
        lin = k["lineage"]
        return [k["name"], CATEGORY_LABELS.get(k["category"], k["category"]), k["value"],
                k["unit"], k["yoy_change"], k["yoy_change_pct"], k["validation_status"],
                lin["extraction_method"], lin["confidence"], lin["document_title"],
                lin["page"], lin["source_text"]]

    _sheet(wb, "Validated KPIs", kpi_headers, [kpi_row(k) for k in kpis])
    for category, label in CATEGORY_LABELS.items():
        cat = [k for k in kpis if k["category"] == category]
        if cat:
            _sheet(wb, label[:31], kpi_headers, [kpi_row(k) for k in cat])

    bench = benchmark_summary(db, fy)
    bench_rows = [
        [b["name"], p["bank_name"], p["value"], b["unit"], p["rank"], p["percentile"]]
        for b in bench["kpis"] for p in b["peers"]
    ]
    _sheet(wb, "Benchmarking",
           ["KPI", "Bank", "Value", "Unit", "Rank", "Percentile"], bench_rows)

    trend_rows = []
    for k in kpis:
        hist = warehouse.history(db, bank.id, k["kpi_code"])
        for point in hist["series"]:
            trend_rows.append([k["name"], point["fiscal_year"], point["value"], k["unit"]])
    _sheet(wb, "Trends", ["KPI", "Fiscal Year", "Value", "Unit"], trend_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
