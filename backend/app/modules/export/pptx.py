"""Board Presentation export — consulting-deck structure with executive
headlines, evidence tables and recommendations, in the Old Money theme."""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from sqlalchemy.orm import Session

from app.models import Bank
from app.modules.benchmarking.engine import benchmark_summary
from app.modules.export import theme
from app.modules.export.theme import MODULE_LABELS, format_value
from app.modules.kpi_warehouse import service as warehouse
from app.modules.narrative.engine import build_narrative

INK = RGBColor.from_string(theme.INK)
CARD = RGBColor.from_string(theme.CARD)
GOLD = RGBColor.from_string(theme.GOLD)
TEXT = RGBColor.from_string(theme.TEXT_PRIMARY)
TEXT_2 = RGBColor.from_string(theme.TEXT_SECONDARY)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)  # rectangle
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def _text(slide, left, top, width, height, text, *, size=14, color=TEXT,
          bold=False, font="Georgia"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return box


def _bullets(slide, left, top, width, height, items, *, size=13, color=TEXT_2):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "—  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
        p.space_after = Pt(6)
    return box


def _headline(slide, text):
    _text(slide, Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.0),
          text, size=24, color=GOLD, bold=True)
    rule = slide.shapes.add_shape(1, Inches(0.6), Inches(1.25), Inches(12.1), Pt(1.5))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor.from_string(theme.BORDER)
    rule.line.fill.background()
    rule.shadow.inherit = False


def build_pptx(db: Session, bank: Bank, fy: str) -> bytes:
    narrative = build_narrative(db, bank, fy)
    payload = warehouse.kpi_payload(db, bank, fy)
    kpis = {k["kpi_code"]: k for k in payload["kpis"]}

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    # 1 — Title
    slide = _blank_slide(prs)
    _text(slide, Inches(0.8), Inches(2.3), Inches(11.7), Inches(1.2),
          "SOVEREIGN", size=44, color=GOLD, bold=True)
    _text(slide, Inches(0.8), Inches(3.3), Inches(11.7), Inches(0.8),
          f"{bank.name} — Annual Banking Intelligence Review, {fy}", size=22, color=TEXT)
    _text(slide, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5),
          "Prepared by the Sovereign Banking Intelligence Platform · Single-source KPI warehouse",
          size=11, color=TEXT_2, font="Calibri")

    # 2 — Executive summary
    slide = _blank_slide(prs)
    _headline(slide, "Executive Summary")
    _bullets(slide, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.0),
             narrative["executive_summary"] or ["No summary available."], size=16, color=TEXT)

    # 3 — Key financial highlights (evidence table)
    slide = _blank_slide(prs)
    _headline(slide, "Key Financial Highlights")
    highlight_codes = ["pat", "roe", "roa", "nim", "cost_to_income",
                       "gnpa_ratio", "pcr", "cet1_ratio", "crar", "casa_ratio"]
    rows = [k for c in highlight_codes if (k := kpis.get(c))]
    table_shape = slide.shapes.add_table(
        len(rows) + 1, 4, Inches(0.8), Inches(1.6), Inches(11.7),
        Inches(0.4 * (len(rows) + 1))).table
    for j, header in enumerate(("KPI", "Value", "YoY Δ", "Validation")):
        cell = table_shape.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.color.rgb = GOLD
        para.runs[0].font.size = Pt(13)
        para.runs[0].font.bold = True
    for i, k in enumerate(rows, start=1):
        yoy = f"{k['yoy_change']:+,.2f}" if k["yoy_change"] is not None else "—"
        for j, value in enumerate((k["name"], format_value(k["value"], k["unit"]),
                                   yoy, k["validation_status"])):
            cell = table_shape.cell(i, j)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.color.rgb = TEXT
            para.runs[0].font.size = Pt(12)

    # 4..n — One slide per narrative module
    for section in narrative["sections"]:
        slide = _blank_slide(prs)
        _headline(slide, MODULE_LABELS.get(section["module"], section["module"]))
        _text(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.9),
              section["headline"], size=17, color=TEXT, bold=True)
        _text(slide, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.4),
              section["commentary"], size=13, color=TEXT_2, font="Calibri")
        content = [("Takeaway — " + t) for t in section["takeaways"]] + \
                  [("Recommendation — " + r) for r in section["recommendations"]]
        if content:
            _bullets(slide, Inches(0.8), Inches(4.1), Inches(11.7), Inches(2.9), content)

    # Benchmarking appendix
    bench = benchmark_summary(db, fy)
    slide = _blank_slide(prs)
    _headline(slide, "Appendix — Peer Benchmarking")
    lines = []
    for b in bench["kpis"][:10]:
        me = next((p for p in b["peers"] if p["bank_id"] == bank.id), None)
        if me:
            lines.append(f"{b['name']}: {format_value(me['value'], b['unit'])} — "
                         f"rank {me['rank']} of {len(b['peers'])}")
    _bullets(slide, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.2), lines, size=14)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
